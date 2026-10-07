import json
import re
from collections import defaultdict
from decimal import Decimal, InvalidOperation
from hashlib import sha256
from pathlib import Path

import frappe
from frappe import _
from frappe.model.document import bulk_insert
from frappe.utils import get_datetime, getdate, now_datetime
from openpyxl import load_workbook


BATCH_DOCTYPE = "Temu Financial Reconciliation Batch"

PENDING_DOCTYPE = "Temu Pending Settlement Detail"
TRANSACTION_DOCTYPE = "Temu Transaction Settlement Detail"
PROTECTION_DOCTYPE = "Temu Fulfillment Protection Detail"
LEDGER_DOCTYPE = "Temu Accounting Ledger Detail"
AD_STORE_DOCTYPE = "Temu Advertising Store Daily"
AD_PRODUCT_DOCTYPE = "Temu Advertising Product Report"
AD_RECONCILIATION_DOCTYPE = "Temu Advertising Reconciliation Daily"
AD_PAYMENT_DOCTYPE = "Temu Advertising Payment Transaction"
SNAPSHOT_DOCTYPE = "Temu Financial Source JSON Snapshot"

EXTRACTED_DOCTYPES = (
	PENDING_DOCTYPE,
	TRANSACTION_DOCTYPE,
	PROTECTION_DOCTYPE,
	LEDGER_DOCTYPE,
	AD_STORE_DOCTYPE,
	AD_PRODUCT_DOCTYPE,
	AD_RECONCILIATION_DOCTYPE,
	AD_PAYMENT_DOCTYPE,
)

REQUIRED_BUSINESS_FIELDS = {
	PENDING_DOCTYPE: ("stock_order_number",),
	TRANSACTION_DOCTYPE: ("transaction_type", "currency", "accounting_time"),
	PROTECTION_DOCTYPE: ("protection_type", "business_number"),
	LEDGER_DOCTYPE: ("accounting_time", "accounting_type", "currency"),
	AD_STORE_DOCTYPE: ("record_type",),
	AD_PRODUCT_DOCTYPE: ("record_type",),
	AD_RECONCILIATION_DOCTYPE: ("report_date",),
	AD_PAYMENT_DOCTYPE: ("payment_time", "settlement_number", "payment_status"),
}

REGIONS = {
	"欧区待处理": "欧区",
	"全球待处理": "全球",
	"美区待处理": "美区",
	"欧区财务": "欧区",
	"全球财务": "全球",
	"美区财务": "美区",
}

PENDING_FILE_TYPES = {"欧区待处理", "全球待处理", "美区待处理"}
FINANCE_FILE_TYPES = {"欧区财务", "全球财务", "美区财务"}

AD_STORE_FILE_TYPE = "广告》数据报表》店铺数据报表"
AD_PRODUCT_FILE_TYPE = "广告》数据报表》商品数据报表"
AD_RECONCILIATION_FILE_TYPE = "广告》财务管理》推广流水》对账单"
AD_PAYMENT_FILE_TYPE = "广告》财务管理》推广流水》已支付"


def _text(value):
	if value is None:
		return None
	if isinstance(value, float) and value.is_integer():
		return str(int(value))
	value = str(value).strip()
	return None if value in {"", "/"} else value


def _number(value):
	value = _text(value)
	if value is None or "∞" in value:
		return None
	cleaned = re.sub(r"[¥￥,$,%\s]", "", value)
	if cleaned.startswith("(") and cleaned.endswith(")"):
		cleaned = f"-{cleaned[1:-1]}"
	try:
		return float(Decimal(cleaned))
	except (InvalidOperation, ValueError):
		frappe.throw(_("无法识别数值：{0}").format(value))


def _integer(value):
	number = _number(value)
	return int(number) if number is not None else None


def _is_infinite(value):
	return 1 if "∞" in (_text(value) or "") else 0


def _datetime(value):
	value = _text(value)
	if not value:
		return None
	value = re.sub(r"\s+(CST|UTC|GMT)$", "", value, flags=re.IGNORECASE)
	return get_datetime(value)


def _date(value):
	value = _text(value)
	if not value or value.startswith("共"):
		return None
	return getdate(value)


def _raw_row(values):
	return json.dumps(list(values), ensure_ascii=False, default=str)


def _workbook_json(batch, file_row, file_doc, workbook):
	sheets = []
	for sheet in workbook.worksheets:
		rows = []
		for row_number, values in enumerate(sheet.iter_rows(values_only=True), start=1):
			rows.append({"row_number": row_number, "values": list(values)})
		sheets.append(
			{
				"name": sheet.title,
				"max_row": sheet.max_row,
				"max_column": sheet.max_column,
				"rows": rows,
			}
		)
	payload = {
		"format_version": 1,
		"batch": batch.name,
		"file": {
			"record": file_row.name,
			"type": file_row.file_type,
			"name": file_row.original_file_name or file_doc.file_name,
			"url": file_row.file,
			"hash": file_doc.content_hash or file_row.file_hash,
			"region": REGIONS.get(file_row.file_type),
		},
		"sheets": sheets,
	}
	json_text = json.dumps(
		payload,
		ensure_ascii=False,
		default=str,
		separators=(",", ":"),
	)
	return payload, json_text


def _has_values(values):
	return any(_text(value) is not None for value in values)


def _iter_rows(sheet, start_row):
	for row_number, values in enumerate(
		sheet.iter_rows(min_row=start_row, values_only=True), start=start_row
	):
		values = tuple(values)
		if _has_values(values):
			yield row_number, values


def _header(sheet, row_number=1):
	values = next(
		sheet.iter_rows(min_row=row_number, max_row=row_number, values_only=True), ()
	)
	return tuple(_text(value) for value in values)


def _require_header(sheet, expected, row_number=1):
	actual = list(_header(sheet, row_number))
	while actual and actual[-1] is None:
		actual.pop()
	actual = tuple(actual)
	if actual != tuple(expected):
		frappe.throw(
			_(
				"工作表 {0} 的第 {1} 行表头与预期格式不一致。"
				"可能存在新增、删除、改名或调整顺序的列。"
			).format(sheet.title, row_number)
		)


def _validate_business_record(doctype, data, file_row, sheet, row_number):
	missing = [
		fieldname
		for fieldname in REQUIRED_BUSINESS_FIELDS[doctype]
		if data.get(fieldname) in (None, "")
	]
	if missing:
		frappe.throw(
			_("文件 {0}、工作表 {1}、第 {2} 行缺少关键数据：{3}").format(
				file_row.original_file_name or file_row.file_type,
				sheet.title,
				row_number,
				", ".join(missing),
			)
		)


def _file_document(batch_name, file_url):
	file_name = frappe.db.get_value(
		"File",
		{
			"attached_to_doctype": BATCH_DOCTYPE,
			"attached_to_name": batch_name,
			"file_url": file_url,
		},
		"name",
		order_by="creation desc",
	)
	if not file_name:
		frappe.throw(_("找不到批次 {0} 的文件：{1}").format(batch_name, file_url))
	return frappe.get_doc("File", file_name)


def _source_values(batch, file_row, file_doc, sheet, row_number, values, region=None):
	raw_json = _raw_row(values)
	hash_payload = {
		"batch": batch.name,
		"file_record": file_row.name,
		"file_type": file_row.file_type,
		"file_name": file_row.original_file_name or file_doc.file_name,
		"file_hash": file_doc.content_hash or file_row.file_hash,
		"sheet": sheet.title,
		"row": row_number,
		"raw": raw_json,
	}
	row_hash = sha256(
		json.dumps(hash_payload, ensure_ascii=False, sort_keys=True).encode("utf-8")
	).hexdigest()
	return {
		"name": row_hash[:16],
		"reconciliation_batch": batch.name,
		"source_file_record": file_row.name,
		"source_file_type": file_row.file_type,
		"region": region,
		"source_file_name": file_row.original_file_name or file_doc.file_name,
		"source_sheet": sheet.title,
		"source_row_number": row_number,
		"source_file_hash": file_doc.content_hash or file_row.file_hash,
		"source_row_hash": row_hash,
		"imported_at": now_datetime(),
		"raw_row_json": raw_json,
	}


def _add_record(
	result, doctype, batch, file_row, file_doc, sheet, row_number, values, data, region=None
):
	_validate_business_record(doctype, data, file_row, sheet, row_number)
	record = _source_values(
		batch, file_row, file_doc, sheet, row_number, values, region=region
	)
	record.update(data)
	result[doctype].append(record)


def _parse_pending(result, batch, file_row, file_doc, sheet):
	_require_header(
		sheet,
		(
			"备货单号",
			"SKU ID",
			"SKU货号",
			"货品名称",
			"SKU属性",
			"销售数量",
			"单品券优惠金额",
			"店铺满减券优惠金额",
			"申报价格折扣金额",
			"预估待结算销售额",
			"币种",
		),
	)
	for row_number, values in _iter_rows(sheet, 2):
		_add_record(
			result,
			PENDING_DOCTYPE,
			batch,
			file_row,
			file_doc,
			sheet,
			row_number,
			values,
			{
				"stock_order_number": _text(values[0]),
				"sku_id": _text(values[1]),
				"sku_code": _text(values[2]),
				"product_name": _text(values[3]),
				"sku_attributes": _text(values[4]),
				"sales_quantity": _number(values[5]),
				"item_coupon_discount": _number(values[6]),
				"shop_coupon_discount": _number(values[7]),
				"declared_price_discount": _number(values[8]),
				"estimated_pending_sales": _number(values[9]),
				"currency": _text(values[10]),
			},
			region=REGIONS[file_row.file_type],
		)


def _parse_transaction(result, batch, file_row, file_doc, sheet):
	_require_header(
		sheet,
		(
			"销售订单",
			"售后订单",
			"备货单号",
			"备货单类型",
			"SKU ID",
			"SKU货号",
			"货品名称",
			"SKU属性",
			"数量",
			"单品券金额",
			"店铺满减券金额",
			"申报价格折扣金额",
			"交易类型",
			"金额",
			"币种",
			"账务时间",
		),
	)
	for row_number, values in _iter_rows(sheet, 2):
		_add_record(
			result,
			TRANSACTION_DOCTYPE,
			batch,
			file_row,
			file_doc,
			sheet,
			row_number,
			values,
			{
				"sales_order_number": _text(values[0]),
				"after_sales_order_number": _text(values[1]),
				"stock_order_number": _text(values[2]),
				"stock_order_type": _text(values[3]),
				"sku_id": _text(values[4]),
				"sku_code": _text(values[5]),
				"product_name": _text(values[6]),
				"sku_attributes": _text(values[7]),
				"quantity": _number(values[8]),
				"item_coupon_amount": _number(values[9]),
				"shop_coupon_amount": _number(values[10]),
				"declared_price_discount_amount": _number(values[11]),
				"transaction_type": _text(values[12]),
				"amount": _number(values[13]),
				"currency": _text(values[14]),
				"accounting_time": _datetime(values[15]),
			},
			region=REGIONS[file_row.file_type],
		)


def _parse_after_sales_issue(result, batch, file_row, file_doc, sheet):
	_require_header(
		sheet,
		("违规ID", "货品名称", "SKU", None, None, "赔付金额", "币种", "账务时间"),
	)
	for row_number, values in _iter_rows(sheet, 3):
		business_number = _text(values[0])
		_add_record(
			result,
			PROTECTION_DOCTYPE,
			batch,
			file_row,
			file_doc,
			sheet,
			row_number,
			values,
			{
				"protection_type": "售后问题",
				"business_number": business_number,
				"violation_id": business_number,
				"product_name": _text(values[1]),
				"sku_id": _text(values[2]),
				"sku_code": _text(values[3]),
				"sku_attributes": _text(values[4]),
				"compensation_amount": _number(values[5]),
				"currency": _text(values[6]),
				"accounting_time": _datetime(values[7]),
			},
			region=REGIONS[file_row.file_type],
		)


def _parse_after_sales_reshipment(result, batch, file_row, file_doc, sheet):
	_require_header(sheet, ("订单编号", "SKU", None, None, None, "数量", "赔付金额"))
	for row_number, values in _iter_rows(sheet, 3):
		business_number = _text(values[0])
		_add_record(
			result,
			PROTECTION_DOCTYPE,
			batch,
			file_row,
			file_doc,
			sheet,
			row_number,
			values,
			{
				"protection_type": "售后补寄",
				"business_number": business_number,
				"order_number": business_number,
				"sku_id": _text(values[1]),
				"product_name": _text(values[2]),
				"sku_code": _text(values[3]),
				"sku_attributes": _text(values[4]),
				"quantity": _number(values[5]),
				"compensation_amount": _number(values[6]),
				"currency": "CNY",
			},
			region=REGIONS[file_row.file_type],
		)


def _parse_ledger(result, batch, file_row, file_doc, sheet):
	_require_header(sheet, ("账务时间", "账务类型", "币种", "收支金额", "备注"))
	for row_number, values in _iter_rows(sheet, 2):
		_add_record(
			result,
			LEDGER_DOCTYPE,
			batch,
			file_row,
			file_doc,
			sheet,
			row_number,
			values,
			{
				"accounting_time": _datetime(values[0]),
				"accounting_type": _text(values[1]),
				"currency": _text(values[2]),
				"income_expense_amount": _number(values[3]),
				"remarks": _text(values[4]),
			},
		)


def _advertising_metrics(values, offset, spend_field, net_spend_field):
	return {
		spend_field: _number(values[offset]),
		net_spend_field: _number(values[offset + 1]),
		"declared_sales_amount": _number(values[offset + 2]),
		"roas": _number(values[offset + 3]),
		"roas_infinite": _is_infinite(values[offset + 3]),
		"expense_ratio": _number(values[offset + 4]),
		"cost_per_order": _number(values[offset + 5]),
		"suborder_count": _integer(values[offset + 6]),
		"units": _integer(values[offset + 7]),
		"impressions": _integer(values[offset + 8]),
		"clicks": _integer(values[offset + 9]),
		"click_through_rate": _number(values[offset + 10]),
		"conversion_rate": _number(values[offset + 11]),
		"add_to_cart_count": _integer(values[offset + 12]),
		"net_declared_sales_amount": _number(values[offset + 13]),
		"net_roas": _number(values[offset + 14]),
		"net_roas_infinite": _is_infinite(values[offset + 14]),
		"net_expense_ratio": _number(values[offset + 15]),
		"net_cost_per_order": _number(values[offset + 16]),
		"net_suborder_count": _integer(values[offset + 17]),
		"net_units": _integer(values[offset + 18]),
	}


def _parse_ad_store(result, batch, file_row, file_doc, sheet):
	_require_header(
		sheet,
		(
			"日期",
			"总花费",
			"净总花费",
			"申报价销售额（全域）",
			"投资回报率(ROAS)（全域）",
			"费比（全域）",
			"每笔成交花费（全域）",
			"子订单数（全域）",
			"件数（全域）",
			"曝光（全域）",
			"点击（全域）",
			"点击率（全域）",
			"转化率（全域）",
			"加入购物车数（全域）",
			"净申报价销售额（全域）",
			"净投资回报率(ROAS)（全域）",
			"净费比（全域）",
			"净每笔成交花费（全域）",
			"净子订单数（全域）",
			"净件数（全域）",
		),
	)
	for row_number, values in _iter_rows(sheet, 2):
		data = {
			"record_type": "汇总" if (_text(values[0]) or "").startswith("共") else "每日",
			"report_date": _date(values[0]),
			"currency": "CNY",
		}
		data.update(_advertising_metrics(values, 1, "total_spend", "net_total_spend"))
		_add_record(
			result,
			AD_STORE_DOCTYPE,
			batch,
			file_row,
			file_doc,
			sheet,
			row_number,
			values,
			data,
		)


def _parse_ad_product(result, batch, file_row, file_doc, sheet):
	_require_header(
		sheet,
		(
			"商品名称",
			"商品ID",
			"SPU ID",
			"花费",
			"净花费",
			"申报价销售额（全域）",
			"投资回报率(ROAS)（全域）",
			"费比（全域）",
			"每笔成交花费（全域）",
			"子订单数（全域）",
			"件数（全域）",
			"曝光（全域）",
			"点击（全域）",
			"点击率（全域）",
			"转化率（全域）",
			"加入购物车数（全域）",
			"净申报价销售额（全域）",
			"净投资回报率(ROAS)（全域）",
			"净费比（全域）",
			"净每笔成交花费（全域）",
			"净子订单数（全域）",
			"净件数（全域）",
		),
	)
	for row_number, values in _iter_rows(sheet, 2):
		is_summary = (_text(values[0]) or "").startswith("共")
		data = {
			"record_type": "汇总" if is_summary else "商品",
			"period_start": getdate(batch.start_date) if batch.start_date else None,
			"period_end": getdate(batch.end_date) if batch.end_date else None,
			"product_name": None if is_summary else _text(values[0]),
			"product_id": _text(values[1]),
			"spu_id": _text(values[2]),
			"currency": "CNY",
		}
		data.update(_advertising_metrics(values, 3, "spend", "net_spend"))
		_add_record(
			result,
			AD_PRODUCT_DOCTYPE,
			batch,
			file_row,
			file_doc,
			sheet,
			row_number,
			values,
			data,
		)


def _parse_ad_reconciliation(result, batch, file_row, file_doc, sheet):
	_require_header(
		sheet,
		("日期", "本期花费", "本期已付款金额", "本期优惠金额", "本期已使用退单红包金额"),
	)
	for row_number, values in _iter_rows(sheet, 2):
		_add_record(
			result,
			AD_RECONCILIATION_DOCTYPE,
			batch,
			file_row,
			file_doc,
			sheet,
			row_number,
			values,
			{
				"report_date": _date(values[0]),
				"currency": "CNY",
				"period_spend": _number(values[1]),
				"period_paid_amount": _number(values[2]),
				"period_discount_amount": _number(values[3]),
				"period_refund_red_packet_amount": _number(values[4]),
			},
		)


def _parse_ad_payment(result, batch, file_row, file_doc, sheet):
	_require_header(
		sheet,
		("付款时间", "结算单生成时间", "结算单号", "站点", "结算方式", "流水类型", "交易金额", "操作人", "状态"),
	)
	for row_number, values in _iter_rows(sheet, 2):
		_add_record(
			result,
			AD_PAYMENT_DOCTYPE,
			batch,
			file_row,
			file_doc,
			sheet,
			row_number,
			values,
			{
				"payment_time": _datetime(values[0]),
				"settlement_generated_time": _datetime(values[1]),
				"settlement_number": _text(values[2]),
				"site": _text(values[3]),
				"settlement_method": _text(values[4]),
				"flow_type": _text(values[5]),
				"currency": "CNY",
				"transaction_amount": _number(values[6]),
				"operator": _text(values[7]),
				"payment_status": _text(values[8]),
			},
		)


def _parse_sheet(result, batch, file_row, file_doc, sheet):
	file_type = file_row.file_type
	if file_type in PENDING_FILE_TYPES:
		if sheet.title != "交易结算":
			frappe.throw(_("文件类型 {0} 包含无法识别的工作表：{1}").format(file_type, sheet.title))
		_parse_pending(result, batch, file_row, file_doc, sheet)
		return

	if file_type in FINANCE_FILE_TYPES:
		if sheet.title == "交易结算":
			_parse_transaction(result, batch, file_row, file_doc, sheet)
		elif sheet.title == "消费者及履约保障-售后问题":
			_parse_after_sales_issue(result, batch, file_row, file_doc, sheet)
		elif sheet.title == "消费者及履约保障-售后补寄":
			_parse_after_sales_reshipment(result, batch, file_row, file_doc, sheet)
		else:
			frappe.throw(_("区域财务文件包含无法识别的工作表：{0}").format(sheet.title))
		return

	if file_type == "财务明细":
		_parse_ledger(result, batch, file_row, file_doc, sheet)
	elif file_type == AD_STORE_FILE_TYPE:
		_parse_ad_store(result, batch, file_row, file_doc, sheet)
	elif file_type == AD_PRODUCT_FILE_TYPE:
		_parse_ad_product(result, batch, file_row, file_doc, sheet)
	elif file_type == AD_RECONCILIATION_FILE_TYPE:
		_parse_ad_reconciliation(result, batch, file_row, file_doc, sheet)
	elif file_type == AD_PAYMENT_FILE_TYPE:
		_parse_ad_payment(result, batch, file_row, file_doc, sheet)
	else:
		frappe.throw(_("无法识别来源文件类型：{0}").format(file_type))


def _build_expected_records(batch):
	result = defaultdict(list)
	snapshots = []
	for file_row in sorted(batch.files or [], key=lambda row: row.sort_order or row.idx):
		if not file_row.file:
			continue
		file_doc = _file_document(batch.name, file_row.file)
		if not (file_doc.file_name or "").lower().endswith(".xlsx"):
			frappe.throw(_("只支持解析 XLSX 文件：{0}").format(file_doc.file_name))
		path = Path(file_doc.get_full_path())
		if not path.exists():
			frappe.throw(_("文件实体不存在：{0}").format(file_doc.file_name))
		try:
			workbook = load_workbook(path, read_only=True, data_only=True)
		except Exception as error:
			frappe.throw(_("无法读取文件 {0}：{1}").format(file_doc.file_name, error))
		try:
			snapshot_payload, json_text = _workbook_json(
				batch, file_row, file_doc, workbook
			)
			rows_before = sum(len(rows) for rows in result.values())
			for sheet in workbook.worksheets:
				if sheet.max_row == 1 and sheet.max_column == 1 and not _has_values(_header(sheet)):
					continue
				_parse_sheet(result, batch, file_row, file_doc, sheet)
			rows_after = sum(len(rows) for rows in result.values())
			json_hash = sha256(json_text.encode("utf-8")).hexdigest()
			snapshot_key = sha256(
				f"{batch.name}|{file_row.name}".encode("utf-8")
			).hexdigest()
			snapshots.append(
				{
					"name": snapshot_key[:16],
					"reconciliation_batch": batch.name,
					"source_file_record": file_row.name,
					"source_file_type": file_row.file_type,
					"region": REGIONS.get(file_row.file_type),
					"source_file_name": file_row.original_file_name or file_doc.file_name,
					"source_file_url": file_row.file,
					"source_file_hash": file_doc.content_hash or file_row.file_hash,
					"snapshot_key": snapshot_key,
					"format_version": snapshot_payload["format_version"],
					"sheet_count": len(snapshot_payload["sheets"]),
					"data_row_count": rows_after - rows_before,
					"json_size": len(json_text.encode("utf-8")),
					"json_hash": json_hash,
					"workbook_json": json_text,
					"conversion_status": "成功",
					"converted_at": now_datetime(),
				}
			)
		finally:
			workbook.close()

	for doctype in EXTRACTED_DOCTYPES:
		result.setdefault(doctype, [])
	return result, snapshots


def _is_complete(batch_name, expected, snapshots):
	expected_snapshots = {
		row["source_file_record"]: (
			row["source_file_hash"],
			row["json_hash"],
			row["sheet_count"],
			row["data_row_count"],
		)
		for row in snapshots
	}
	existing_snapshots = {
		row.source_file_record: (
			row.source_file_hash,
			row.json_hash,
			row.sheet_count,
			row.data_row_count,
		)
		for row in frappe.get_all(
			SNAPSHOT_DOCTYPE,
			filters={"reconciliation_batch": batch_name},
			fields=[
				"source_file_record",
				"source_file_hash",
				"json_hash",
				"sheet_count",
				"data_row_count",
			],
		)
	}
	if expected_snapshots != existing_snapshots:
		return False

	for doctype in EXTRACTED_DOCTYPES:
		expected_hashes = {row["source_row_hash"] for row in expected[doctype]}
		existing_hashes = set(
			frappe.get_all(
				doctype,
				filters={"reconciliation_batch": batch_name},
				pluck="source_row_hash",
			)
		)
		if expected_hashes != existing_hashes:
			return False
	return True


def delete_batch_data(batch_name):
	frappe.db.delete(SNAPSHOT_DOCTYPE, {"reconciliation_batch": batch_name})
	for doctype in EXTRACTED_DOCTYPES:
		frappe.db.delete(doctype, {"reconciliation_batch": batch_name})


def _bulk_insert_values(doctype, values_list):
	if not values_list:
		return
	documents = [
		frappe.get_doc({"doctype": doctype, **values}) for values in values_list
	]
	bulk_insert(doctype, documents, chunk_size=500)


def _insert_records(expected, snapshots):
	_bulk_insert_values(SNAPSHOT_DOCTYPE, snapshots)
	for doctype in EXTRACTED_DOCTYPES:
		_bulk_insert_values(doctype, expected[doctype])


def sync_batch_data(batch_name, force=False):
	batch = frappe.get_doc(BATCH_DOCTYPE, batch_name)
	expected, snapshots = _build_expected_records(batch)
	counts = {doctype: len(expected[doctype]) for doctype in EXTRACTED_DOCTYPES}
	total = sum(counts.values())

	if not force and _is_complete(batch.name, expected, snapshots):
		if batch.files:
			frappe.db.set_value(BATCH_DOCTYPE, batch.name, "status", "已处理", update_modified=False)
		return {
			"status": "unchanged",
			"total": total,
			"snapshots": len(snapshots),
			"counts": counts,
		}

	delete_batch_data(batch.name)
	_insert_records(expected, snapshots)
	frappe.db.set_value(
		BATCH_DOCTYPE,
		batch.name,
		"status",
		"已处理" if batch.files else "草稿",
		update_modified=False,
	)
	return {
		"status": "synchronized",
		"total": total,
		"snapshots": len(snapshots),
		"counts": counts,
	}
